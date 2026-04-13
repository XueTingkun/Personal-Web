from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_submission_word():
    doc = Document()
    
    # Title
    title = doc.add_heading('Final Individual Project — Personal Job-Seeking Website', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Student Info
    doc.add_heading('Student Information', level=1)
    p = doc.add_paragraph()
    p.add_run('Name: ').bold = True
    p.add_run('薛廷焜 (Tingkun Xue)\n')
    p.add_run('Student ID: ').bold = True
    p.add_run('5563215\n')
    
    # Website Info
    doc.add_heading('Website Submission', level=1)
    p2 = doc.add_paragraph()
    p2.add_run('Website URL: ').bold = True
    p2.add_run('https://XueTingkun.github.io/Personal-Web/\n')
    p2.add_run('GitHub Repository: ').bold = True
    p2.add_run('https://github.com/XueTingkun/Personal-Web\n')
    
    # Project Description Summary
    doc.add_heading('Project Overview', level=1)
    doc.add_paragraph('This website is a professional personal portfolio built using React, Tailwind CSS, and Vite. It includes all 6 required sections as per the project rubric:')
    
    sections = [
        "Home / Headline: Professional positioning as an AI Engineer.",
        "About Me: Educational background from Lingnan University and Hanshan Normal University.",
        "Skills: Technical, business, and tool-based skills (Python, ML, Dify, RAG, etc.).",
        "Projects: Detailed showcases including China Mobile AI Auditing, Q-learning Dungeon Battle, and a CV Group Project.",
        "Resume: Downloadable CV and professional highlights.",
        "Contact: Direct email and social links."
    ]
    
    for section in sections:
        doc.add_paragraph(section, style='List Bullet')

    # Save the document
    file_name = 'TingkunXue_5563215_Individual_Project_Submission.docx'
    doc.save(file_name)
    print(f"Successfully generated {file_name}")

if __name__ == "__main__":
    create_submission_word()
